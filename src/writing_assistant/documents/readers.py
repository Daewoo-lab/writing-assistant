"""Read a brief file (docx / pdf) into plain text."""

from __future__ import annotations

import io
from pathlib import PurePosixPath

import pdfplumber
from docx import Document


def read_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts).strip()


def read_pdf(data: bytes) -> str:
    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text.strip())
    return "\n\n".join(parts).strip()


def read_any(data: bytes, filename: str) -> str:
    ext = PurePosixPath(filename).suffix.lower()
    if ext == ".docx":
        return read_docx(data)
    if ext == ".pdf":
        return read_pdf(data)
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="replace").strip()
    raise ValueError(f"unsupported brief format: {ext!r}")
