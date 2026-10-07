"""Render generated content into office files (docx / pptx)."""

from __future__ import annotations

import io
from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Cm, Pt
from pptx import Presentation
from pptx.util import Pt as PptPt


@dataclass(frozen=True)
class DocMeta:
    title: str = ""
    author: str = ""


def write_docx(text: str, meta: DocMeta) -> bytes:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for section in doc.sections:
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    if meta.title:
        doc.add_heading(meta.title, level=1)

    for block in text.split("\n\n"):
        block = block.strip()
        if block:
            doc.add_paragraph(block)

    if meta.author:
        doc.add_paragraph()
        doc.add_paragraph(meta.author)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def write_pptx(outline: list[dict[str, object]], meta: DocMeta) -> bytes:
    prs = Presentation()
    blank_title = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    if meta.title:
        slide = prs.slides.add_slide(blank_title)
        slide.shapes.title.text = meta.title
        if slide.placeholders[1] is not None and meta.author:
            slide.placeholders[1].text = meta.author

    for item in outline:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = str(item.get("title", ""))

        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = item.get("bullets") or []
        if isinstance(bullets, list):
            for i, bullet in enumerate(bullets):
                para = body.paragraphs[0] if i == 0 else body.add_paragraph()
                para.text = str(bullet)
                para.font.size = PptPt(18)

        notes = item.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = str(notes)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
