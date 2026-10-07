from __future__ import annotations

import io

import pytest
from docx import Document

from writing_assistant.documents.pipeline import DocumentPipeline, WritingSample
from writing_assistant.documents.readers import read_any, read_docx
from writing_assistant.documents.writers import DocMeta, write_docx, write_pptx


def _make_docx(text: str) -> bytes:
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_round_trip() -> None:
    data = _make_docx("Заголовок брифа\nТребование один\nТребование два")
    text = read_docx(data)
    assert "Требование один" in text
    assert "Требование два" in text


def test_read_any_dispatch_txt() -> None:
    assert read_any("привет".encode(), "brief.txt") == "привет"
    with pytest.raises(ValueError):
        read_any(b"x", "brief.rtf")


def test_write_docx_is_openable() -> None:
    out = write_docx("Первый абзац.\n\nВторой абзац.", DocMeta(title="Отчёт", author="Автор"))
    doc = Document(io.BytesIO(out))
    texts = [p.text for p in doc.paragraphs]
    assert "Отчёт" in texts
    assert "Первый абзац." in texts
    assert "Автор" in texts


def test_write_pptx_is_openable() -> None:
    outline = [{"title": "Слайд 1", "bullets": ["раз", "два"], "notes": "заметка"}]
    out = write_pptx(outline, DocMeta(title="Дека", author="Автор"))
    assert out[:2] == b"PK"  # zip/pptx magic


class _StubLLM:
    async def generate_text(self, brief_text, hints, samples):  # type: ignore[no-untyped-def]
        assert "бриф" in brief_text.lower()
        return f"draft for: {brief_text[:20]}"

    async def generate_presentation_outline(self, brief_text, hints, samples):  # type: ignore[no-untyped-def]
        return [{"title": "T", "bullets": ["b"], "notes": ""}]


@pytest.mark.asyncio
async def test_pipeline_docx() -> None:
    pipe = DocumentPipeline(_StubLLM())  # type: ignore[arg-type]
    brief = _make_docx("Это бриф на статью")
    res = await pipe.generate(brief, "brief.docx", {"title": "Z"}, [WritingSample("x", "y")], "docx")
    assert res.filename == "draft.docx"
    assert res.mime.endswith("wordprocessingml.document")
    assert len(res.content) > 0


@pytest.mark.asyncio
async def test_pipeline_pptx() -> None:
    pipe = DocumentPipeline(_StubLLM())  # type: ignore[arg-type]
    brief = _make_docx("Это бриф на презентацию")
    res = await pipe.generate(brief, "brief.docx", {}, [], "pptx")
    assert res.filename == "draft.pptx"
    assert res.content[:2] == b"PK"
